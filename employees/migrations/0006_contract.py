from django.db import migrations, models


def as_int(value):
    if value is None or value == "":
        return 0
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return 0


def as_text(value):
    if value is None:
        return ""
    try:
        # Excel dates are converted to ISO-like text for consistent display.
        if hasattr(value, "strftime"):
            return value.strftime("%Y-%m-%d")
    except Exception:
        pass
    return str(value).strip()


def import_contracts(apps, schema_editor):
    Contract = apps.get_model("employees", "Contract")
    from pathlib import Path
    from openpyxl import load_workbook

    base = Path(__file__).resolve().parent.parent / "data"
    files = [
        ("1000", base / "contract_1000.xlsx"),
        ("bachelor", base / "contract_bachelor_diploma.xlsx"),
        ("secondary", base / "contract_secondary_below.xlsx"),
    ]

    for category, path in files:
        if not path.exists():
            continue
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            continue
        headers = [str(h).strip() if h is not None else "" for h in headers]
        idx = {h: i for i, h in enumerate(headers) if h}

        for row in rows:
            if not row or not row[0]:
                continue
            code = as_text(row[0])
            if not code:
                continue

            if category == "1000":
                values = dict(
                    category=category,
                    code=code,
                    name=as_text(row[1]),
                    monthly_salary=as_int(row[2]),
                    net_salary=as_int(row[2]),
                    bank=as_text(row[3]),
                )
            elif category == "bachelor":
                values = dict(
                    category=category,
                    code=code,
                    name=as_text(row[1]),
                    start_date=as_text(row[2]),
                    days=as_int(row[3]),
                    monthly_salary=as_int(row[4]),
                    total_amount=as_int(row[5]),
                    absence_deduction=as_int(row[6]),
                    net_after_absence=as_int(row[7]),
                    pension_5=as_int(row[8]),
                    net_salary=as_int(row[9]),
                    bank=as_text(row[10]),
                    phone=as_text(row[11]),
                    notes=as_text(row[12]),
                )
            else:
                values = dict(
                    category=category,
                    code=code,
                    name=as_text(row[1]),
                    job_title=as_text(row[2]),
                    start_date=as_text(row[3]),
                    monthly_salary=as_int(row[4]),
                    absence_deduction=as_int(row[5]),
                    net_after_absence=as_int(row[6]),
                    pension_5=as_int(row[7]),
                    net_salary=as_int(row[8]),
                    bank=as_text(row[9]),
                    workplace=as_text(row[10]),
                    notes=as_text(row[11]),
                    phone=as_text(row[12]),
                )
            Contract.objects.update_or_create(category=category, code=code, defaults=values)


def reverse_import(apps, schema_editor):
    Contract = apps.get_model("employees", "Contract")
    Contract.objects.all().delete()


class Migration(migrations.Migration):
    dependencies = [("employees", "0005_employeeinfo_result")]

    operations = [
        migrations.CreateModel(
            name="Contract",
            fields=[
                ("id", models.BigAutoField(auto_created=True, primary_key=True, serialize=False, verbose_name="ID")),
                ("category", models.CharField(choices=[("1000", "عقود الألف درجة"), ("bachelor", "عقود البكالوريوس والدبلوم"), ("secondary", "عقود إعدادية فما دون")], max_length=20)),
                ("code", models.CharField(max_length=30)),
                ("name", models.CharField(max_length=200)),
                ("job_title", models.CharField(blank=True, default="", max_length=200)),
                ("start_date", models.CharField(blank=True, default="", max_length=50)),
                ("days", models.IntegerField(default=0)),
                ("monthly_salary", models.IntegerField(default=0)),
                ("total_amount", models.IntegerField(default=0)),
                ("absence_deduction", models.IntegerField(default=0)),
                ("net_after_absence", models.IntegerField(default=0)),
                ("pension_5", models.IntegerField(default=0)),
                ("net_salary", models.IntegerField(default=0)),
                ("bank", models.CharField(blank=True, default="", max_length=100)),
                ("workplace", models.CharField(blank=True, default="", max_length=200)),
                ("phone", models.CharField(blank=True, default="", max_length=50)),
                ("notes", models.TextField(blank=True, default="")),
            ],
            options={"ordering": ["name"]},
        ),
        migrations.AddConstraint(
            model_name="contract",
            constraint=models.UniqueConstraint(fields=("category", "code"), name="unique_contract_category_code"),
        ),
        migrations.RunPython(import_contracts, reverse_import),
    ]
