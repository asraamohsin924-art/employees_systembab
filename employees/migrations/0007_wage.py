from django.db import migrations, models


def as_int(value):
    if value is None or value == "":
        return 0
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return 0


def as_text(value):
    if value is None:
        return ""
    return str(value).strip()


def import_wages(apps, schema_editor):
    Wage = apps.get_model("employees", "Wage")
    from pathlib import Path
    from openpyxl import load_workbook
    path = Path(__file__).resolve().parent.parent / "data" / "wages.xlsx"
    if not path.exists():
        return
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    next(rows, None)
    for row in rows:
        if not row or not row[0]:
            continue
        code = as_text(row[0])
        if not code:
            continue
        Wage.objects.update_or_create(code=code, defaults={
            "name": as_text(row[1]),
            "days": as_int(row[2]),
            "daily_amount": as_int(row[3]),
            "bonuses": as_int(row[4]),
            "total_amount": as_int(row[5]),
        })


def reverse_wages(apps, schema_editor):
    Wage = apps.get_model("employees", "Wage")
    Wage.objects.all().delete()


class Migration(migrations.Migration):
    dependencies = [("employees", "0006_contract")]
    operations = [
        migrations.CreateModel(
            name="Wage",
            fields=[
                ("id", models.BigAutoField(auto_created=True, primary_key=True, serialize=False, verbose_name="ID")),
                ("code", models.CharField(max_length=30, unique=True)),
                ("name", models.CharField(max_length=200)),
                ("days", models.IntegerField(default=0)),
                ("daily_amount", models.IntegerField(default=0)),
                ("bonuses", models.IntegerField(default=0)),
                ("total_amount", models.IntegerField(default=0)),
            ],
            options={"ordering": ["name"]},
        ),
        migrations.RunPython(import_wages, reverse_wages),
    ]
